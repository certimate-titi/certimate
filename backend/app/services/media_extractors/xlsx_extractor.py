"""XLSX extractor — openpyxl for .xlsx files."""

from __future__ import annotations

import logging
from pathlib import Path

from app.services.media_extractors.base import ExtractionResult

logger = logging.getLogger(__name__)


class XlsxExtractor:

    def extract(self, file_path: str, resource_name: str = "") -> ExtractionResult:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True, data_only=True)
        parts: list[str] = []
        sections: list[dict] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: list[str] = []

            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                # Skip completely empty rows
                if not any(cells):
                    continue
                rows.append(" | ".join(cells))

            if rows:
                # Build markdown table header from first row
                header = rows[0]
                separator = " | ".join(["---"] * len(rows[0].split(" | ")))
                table_text = f"{header}\n{separator}\n" + "\n".join(rows[1:])

                parts.append(f"## {sheet_name}\n{table_text}")
                sections.append({
                    "title": sheet_name,
                    "content": table_text,
                    "page_start": None,
                    "page_end": None,
                    "depth": 1,
                })

        wb.close()

        raw_text = "\n\n".join(parts)

        return ExtractionResult(
            title=resource_name or Path(file_path).stem,
            content_type="xlsx",
            raw_text=raw_text,
            metadata=f"sheets:{len(sections)}",
            sections=sections,
        )
